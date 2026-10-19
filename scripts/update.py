#!/usr/bin/env python3
"""Refresh the Construction Spending Monitor from the Census C30 release.

Fetches the release, rebuilds the DATA and NOTES lines in index.html, runs the
checks from HANDOFF.md and writes data/summary.json. It never touches anything
else in the page. The workflow commits only if this script exits 0 and reports
changed=true.

Exit codes: 0 = updated, or nothing new; 1 = a check or a required source failed.
"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import json
import os
import re
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

C30 = "https://www.census.gov/construction/c30"
URL_TABLE1_XLSX = f"{C30}/xlsx/release.xlsx"
URL_TABLE1_PDF = f"{C30}/pdf/release.pdf"
URL_RELEASE_TEXT = f"{C30}/current/index.html"
URL_SCHEDULE = f"{C30}/release.html"
URL_TIMESERIES = {
    "total": f"{C30}/xlsx/totsatime.xlsx",
    "private": f"{C30}/xlsx/privsatime.xlsx",
    "public": f"{C30}/xlsx/pubsatime.xlsx",
}
FRED_SERIES = {"total": "TTLCONS", "res": "PRRESCONS", "nonres": "PNRESCONS",
               "pub": "TLPBLCONS", "mfg": "PRMFGCONS"}
URL_FRED = "https://fred.stlouisfed.org/data/{}.txt"

USER_AGENT = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
TIMEOUT = 60
RETRY_WAITS = (5, 15, 45)
FRED_TIMEOUT = 15

START_MARK = "// ===== DATA BLOCKS"
END_MARK = "// ===== END DATA BLOCKS ====="
BLOCK_NAMES = ("DATA", "NOTES", "HIST")

FIELDS = ("total", "res", "nonres", "pub", "mfg")  # order of DATA.months columns 1..5
TIE_TOLERANCE = 5          # $ millions
BIG_REVISION_PCT = 5.0
MINUS = "−"
EN_DASH = "–"


class CheckFailed(Exception):
    """A publishing check failed: change nothing, fail the run."""


class SourceError(Exception):
    """A source could not be fetched or read."""


# ---------------------------------------------------------------------------
# Fetching

def fetch(url: str, timeout: int = TIMEOUT, waits=RETRY_WAITS, sleep=time.sleep) -> bytes:
    """GET with a browser User-Agent: one try plus a retry after each wait."""
    import requests
    last = None
    for attempt in range(len(waits) + 1):
        try:
            r = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=timeout)
            r.raise_for_status()
            return r.content
        except requests.RequestException as e:
            last = e
            if attempt < len(waits):
                sleep(waits[attempt])
    raise SourceError(f"{url}: {last}")


def fetch_once(url: str, timeout: int = FRED_TIMEOUT) -> bytes:
    return fetch(url, timeout=timeout, waits=())


# ---------------------------------------------------------------------------
# Months and formatting

def month_add(m: str, n: int) -> str:
    y, mo = map(int, m.split("-"))
    k = y * 12 + (mo - 1) + n
    return f"{k // 12:04d}-{k % 12 + 1:02d}"


def month_abbr(m: str) -> str:
    y, mo = m.split("-")
    return f"{calendar.month_abbr[int(mo)]} {y}"


def month_full(m: str) -> str:
    return calendar.month_name[int(m.split("-")[1])]


def fmt_date(d: dt.date) -> str:
    return f"{calendar.month_abbr[d.month]} {d.day}, {d.year}"


def pct(a: float, b: float) -> float:
    return (a / b - 1) * 100


def fmt_pct(p: float) -> str:
    r = round(p, 1)
    sign = "+" if r > 0 else MINUS if r < 0 else ""
    return f"{sign}{abs(r):.1f}%"


def fmt_b(millions: float) -> str:
    return f"${millions / 1000:,.1f}B"


def fmt_signed_b(millions: float) -> str:
    r = round(millions / 1000, 1)
    sign = "+" if r > 0 else MINUS if r < 0 else ""
    return f"{sign}${abs(r):,.1f}B"


# ---------------------------------------------------------------------------
# The page: three one-line data blocks

def dump(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"))


def block_lines(lines: list[str]) -> dict[str, int]:
    """Index of the DATA, NOTES and HIST lines (each exactly one line between the markers)."""
    starts = [i for i, l in enumerate(lines) if l.startswith(START_MARK)]
    ends = [i for i, l in enumerate(lines) if l.rstrip("\r\n") == END_MARK]
    if len(starts) != 1 or len(ends) != 1 or ends[0] - starts[0] != len(BLOCK_NAMES) + 1:
        raise CheckFailed("data block markers not found, or not exactly three lines between them")
    idx = {}
    for i in range(starts[0] + 1, ends[0]):
        m = re.match(r"const (\w+) = ", lines[i])
        if not m or m.group(1) not in BLOCK_NAMES or m.group(1) in idx:
            raise CheckFailed(f"unexpected line {i + 1} between the data markers")
        idx[m.group(1)] = i
    return idx


def parse_block(line: str):
    body = line.rstrip("\r\n")
    if not body.endswith(";"):
        raise CheckFailed("data line does not end with ';'")
    return json.loads(body[body.index("=") + 1:-1].strip())


def load_page(html: str):
    lines = html.splitlines(keepends=True)
    idx = block_lines(lines)
    return parse_block(lines[idx["DATA"]]), parse_block(lines[idx["NOTES"]]), lines[idx["HIST"]]


def render_page(html: str, data: dict, notes: list) -> str:
    """Replace only the DATA and NOTES lines. HIST is never rewritten."""
    lines = html.splitlines(keepends=True)
    idx = block_lines(lines)
    for name, obj in (("DATA", data), ("NOTES", notes)):
        old = lines[idx[name]]
        ending = old[len(old.rstrip("\r\n")):]
        lines[idx[name]] = f"const {name} = {dump(obj)};{ending}"
    return "".join(lines)


def extract_script(html: str) -> str:
    scripts = re.findall(r"<script\b[^>]*>(.*?)</script>", html, flags=re.S | re.I)
    if len(scripts) != 1:
        raise CheckFailed(f"expected one <script>, found {len(scripts)}")
    return scripts[0]


def node_check(js: str) -> None:
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
        f.write(js)
        path = f.name
    try:
        r = subprocess.run(["node", "--check", path], capture_output=True, text=True)
    except FileNotFoundError:
        raise CheckFailed("node not found; cannot check the page script")
    finally:
        os.unlink(path)
    if r.returncode != 0:
        raise CheckFailed("page script does not parse: " + (r.stderr.strip().splitlines() or ["?"])[-1])


def verify_page(before: str, after: str) -> None:
    """Checks 5, 6 and 8."""
    b, a = before.splitlines(keepends=True), after.splitlines(keepends=True)
    ib, ia = block_lines(b), block_lines(a)
    if ib != ia or len(a) != len(b):
        raise CheckFailed("check 5: page line structure changed")
    skip = set(ib.values())
    diff = [i + 1 for i in range(len(b)) if i not in skip and a[i] != b[i]]
    if diff:
        raise CheckFailed(f"check 5: lines outside the data blocks changed: {diff[:10]}")
    if a[ia["HIST"]] != b[ib["HIST"]]:
        raise CheckFailed("check 6: HIST changed")
    node_check(extract_script(after))


# ---------------------------------------------------------------------------
# Table 1 (release.xlsx, fallback release.pdf)

@dataclass
class Table1:
    cols: list[str]                 # month of each value column: 5 recent, newest first, then year-ago
    values: dict[str, list]         # key -> one value per column (None where not needed/unknown)
    source: str = ""

    @property
    def month(self) -> str:
        return self.cols[0]

    @property
    def yago(self) -> str:
        return self.cols[-1]

    def get(self, key: str, month: str):
        return self.values[key][self.cols.index(month)]


HEADINGS = {"Total Construction": "all", "Total Private Construction": "private",
            "Total Public Construction": "public"}
# (block, exact label) -> key
TABLE1_ROWS = {
    ("all", "Total Construction"): "total",
    ("private", "Total Private Construction"): "private_total",
    ("private", "Residential"): "res",
    ("private", "Nonresidential"): "nonres",
    ("private", "Manufacturing"): "mfg",
    ("private", "Power"): "power",
    ("private", "Office"): "office",
    ("private", "Commercial"): "commercial",
    ("public", "Total Public Construction"): "pub",
    ("public", "Highway and street"): "highway",
    ("public", "Educational"): "educational",
}
FOOTNOTED = {"Total Private Construction", "Total Public Construction", "Total Construction", "Residential"}


def clean_label(label: str) -> str:
    """Trim spaces and drop a footnote digit glued to the four labels that carry one."""
    s = " ".join(str(label).split())
    m = re.fullmatch(r"(.*[A-Za-z])\d", s)
    if m and m.group(1) in FOOTNOTED:
        return m.group(1)
    return s


def table1_from_rows(cols: list[str], rows: list[tuple[str, list]], source: str = "") -> Table1:
    """Map Table 1 rows to keys: split into blocks at the three headings, then match exact labels.

    `rows` is [(label, [one value per column]), ...] in table order, as a reader extracts them.
    """
    if len(cols) != 6:
        raise SourceError(f"Table 1: expected 6 value columns, got {len(cols)}")
    block = None
    found: dict[str, list] = {}
    for label, vals in rows:
        lab = clean_label(label)
        if lab in HEADINGS:
            block = HEADINGS[lab]
        key = TABLE1_ROWS.get((block, lab))
        if key is None:
            continue
        if key in found:
            raise SourceError(f"Table 1: label {lab!r} appears twice in the {block} block")
        if len(vals) != 6:
            raise SourceError(f"Table 1: row {lab!r} has {len(vals)} values, expected 6")
        found[key] = list(vals)
    missing = sorted(set(TABLE1_ROWS.values()) - set(found))
    if missing:
        raise SourceError(f"Table 1: rows not found: {missing}")
    for k in ("total", "private_total", "res", "nonres", "mfg", "pub"):
        if any(not isinstance(v, int) for v in found[k][:5]):
            raise SourceError(f"Table 1: row {k} has non-integer values in the recent months")
    return Table1(cols=list(cols), values=found, source=source)


TABLE1_TITLE = "Table 1. Value of Construction Put in Place in the United States, Seasonally Adjusted Annual Rate"
MONTH_ABBR = {a: i for i, a in enumerate(calendar.month_abbr) if a}
MONTH_NAME = {n: i for i, n in enumerate(calendar.month_name) if n}


def table1_columns(heads: list[str]) -> list[str]:
    """['Aug 2026p', 'Jul 2026r', ...] -> month ids; checks 5 consecutive months plus the year-ago month."""
    cols = []
    for h in heads:
        m = re.fullmatch(r"([A-Z][a-z]{2}) (\d{4})([pr]?)", " ".join(str(h).split()))
        if not m or m.group(1) not in MONTH_ABBR:
            raise SourceError(f"Table 1: unexpected column heading {h!r}")
        cols.append(f"{m.group(2)}-{MONTH_ABBR[m.group(1)]:02d}")
    if len(cols) != 6 or cols[1:5] != [month_add(cols[0], -i) for i in range(1, 5)] or cols[5] != month_add(cols[0], -12):
        raise SourceError(f"Table 1: unexpected column months {cols}")
    return cols


def table1_from_sheet(rows: list[list]) -> Table1:
    """release.xlsx sheet 'Table1': title in A1, headings in the row whose A cell is 'Type of Construction'
    (cells like 'Aug\\n2026p' in B..G, percent changes in H..I), then one row per category."""
    if not rows or str(rows[0][0]).strip() != TABLE1_TITLE:
        raise SourceError("release.xlsx: Table1 title not found in A1")
    hi = next((i for i, r in enumerate(rows) if r and " ".join(str(r[0]).split()) == "Type of Construction"), None)
    if hi is None:
        raise SourceError("release.xlsx: 'Type of Construction' heading row not found")
    cols = table1_columns(rows[hi][1:7])
    out = []
    for r in rows[hi + 1:]:
        vals = list(r[1:7]) + [None] * (6 - len(r[1:7]))
        if isinstance(r[0], str) and all(isinstance(v, int) for v in vals):
            out.append((r[0], vals))
    return table1_from_rows(cols, out, "release.xlsx")


def read_table1_xlsx(content: bytes) -> Table1:
    import io
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb["Table1"]
    except Exception as e:  # not a workbook, or no Table1 sheet
        raise SourceError(f"release.xlsx: {e}")
    return table1_from_sheet([list(r) for r in ws.iter_rows(values_only=True)])


PDF_ROW = re.compile(r"(?P<label>\S.*?) (?P<vals>(?:\d{1,3}(?:,\d{3})* ){5}\d{1,3}(?:,\d{3})*) \S+ \S+")


def table1_from_pdf_text(text: str) -> Table1:
    """pdfplumber text of the Table 1 page: headings split over two lines
    ('Type of Construction Aug Jul ...' then '2026p 2026r ...'), then 'Label 2,203,129 ... 0.9 -1.7'."""
    lines = [" ".join(l.split()) for l in text.splitlines()]
    if not lines or lines[0] != TABLE1_TITLE:
        raise SourceError("release.pdf: Table 1 title not found")
    hi = next((i for i, l in enumerate(lines) if l.startswith("Type of Construction ")), None)
    if hi is None or hi + 1 >= len(lines):
        raise SourceError("release.pdf: heading line not found")
    names = lines[hi].split()[3:9]
    years = lines[hi + 1].split()[:6]
    cols = table1_columns([f"{n} {y}" for n, y in zip(names, years)])
    rows = []
    for l in lines[hi + 2:]:
        if l.startswith("p Preliminary"):
            break
        m = PDF_ROW.fullmatch(l)
        if m:
            rows.append((m.group("label"), [int(v.replace(",", "")) for v in m.group("vals").split()]))
    return table1_from_rows(cols, rows, "release.pdf")


def read_table1_pdf(content: bytes) -> Table1:
    import io
    import pdfplumber
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            texts = [p.extract_text() or "" for p in pdf.pages]
    except Exception as e:
        raise SourceError(f"release.pdf: {e}")
    page = next((t for t in texts if t.startswith(TABLE1_TITLE)), None)
    if page is None:
        raise SourceError("release.pdf: no page starts with the Table 1 title")
    return table1_from_pdf_text(page)


# ---------------------------------------------------------------------------
# Release text, schedule, time series, FRED

@dataclass
class Change:
    level_b: float          # $ billions as printed
    pct: float              # signed percent change as printed
    significant: bool       # False when the asterisk is present


@dataclass
class ReleaseText:
    month: str              # reference month, YYYY-MM
    released: dt.date
    total: Change           # m/m
    yoy: Change             # total vs year-ago
    ytd_months: int
    ytd_cur_b: float
    ytd_prior_b: float
    res: Change
    nonres: Change
    pub: Change


def visible_text(html: bytes | str) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return " ".join(soup.get_text(" ").split())


NUMBER_WORDS = {w: i for i, w in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve".split())}
B = r"\$([\d,]+\.\d) billion"
# "0.9 percent (±1.0 percent)* above": the asterisk (plain '*') marks a change that is not significant
CHG = r"(\d+\.\d) percent \(±(\d+\.\d) percent\)\s?(\*?)\s?(above|below)"


def _change(level: str, p: str, star: str, direction: str) -> Change:
    v = float(p)
    return Change(float(level.replace(",", "")), v if direction == "above" else -v, star != "*")


def _one(pattern: str, text: str, what: str) -> re.Match:
    found = list(re.finditer(pattern, text))
    if len(found) != 1:
        raise SourceError(f"release text: {what} sentence found {len(found)} times")
    return found[0]


def parse_release_text(html: bytes | str) -> ReleaseText:
    """The current release page, read as visible text. Sentences are matched only inside their section."""
    text = visible_text(html)
    rel = re.search(r"FOR IMMEDIATE RELEASE: [A-Z][a-z]+day, ([A-Z][a-z]+ \d{1,2}, \d{4})", text)
    title = re.search(r"Monthly Construction Spending, ([A-Z][a-z]+) (\d{4})", text)
    if not rel or not title or title.group(1) not in MONTH_NAME:
        raise SourceError("release text: release line or page heading not found")
    released = dt.datetime.strptime(rel.group(1), "%B %d, %Y").date()
    mon, year = title.group(1), int(title.group(2))
    month = f"{year}-{MONTH_NAME[mon]:02d}"
    prev_mon = calendar.month_name[int(month_add(month, -1)[5:])]

    i_tot = text.find("Total Construction", title.end())
    i_priv = text.find("Private Construction", i_tot + 1)
    i_pub = text.find("Public Construction", i_priv + 1)
    if min(i_tot, i_priv, i_pub) < 0:
        raise SourceError("release text: section headings not found")
    i_end = text.find("Additional Release Materials", i_pub)
    total_s, priv_s, pub_s = text[i_tot:i_priv], text[i_priv:i_pub], text[i_pub:i_end if i_end > 0 else None]

    m = _one(rf"Construction spending during {mon} {year} was estimated at a seasonally adjusted annual rate "
             rf"of {B}, {CHG} the revised {prev_mon} estimate of {B}\.", total_s, "total")
    total = _change(m.group(1), m.group(2), m.group(4), m.group(5))
    m = _one(rf"The {mon} figure is {CHG} the {mon} {year - 1} estimate of {B}\.", total_s, "year-over-year")
    yoy = _change(m.group(5), m.group(1), m.group(3), m.group(4))
    m = _one(rf"During the first ([a-z]+) months of this year, construction spending amounted to {B}, {CHG} "
             rf"the {B} for the same period in {year - 1}\.", total_s, "year-to-date")
    if m.group(1) not in NUMBER_WORDS:
        raise SourceError(f"release text: unknown month count {m.group(1)!r}")
    ytd_months, ytd_cur, ytd_prior = NUMBER_WORDS[m.group(1)], float(m.group(2).replace(",", "")), float(m.group(7).replace(",", ""))

    def private(kind):
        m = _one(rf"(?<![A-Za-z]){kind} construction was at a seasonally adjusted annual rate of {B} in {mon}, "
                 rf"{CHG} the revised", priv_s, f"private {kind.lower()}")
        return _change(m.group(1), m.group(2), m.group(4), m.group(5))
    m = _one(rf"In {mon}, the estimated seasonally adjusted annual rate of public construction spending was {B}, "
             rf"{CHG} the revised", pub_s, "public")
    pub = _change(m.group(1), m.group(2), m.group(4), m.group(5))
    return ReleaseText(month=month, released=released, total=total, yoy=yoy, ytd_months=ytd_months,
                       ytd_cur_b=ytd_cur, ytd_prior_b=ytd_prior, res=private("Residential"),
                       nonres=private("Nonresidential"), pub=pub)


def parse_schedule(html: bytes | str) -> list[tuple[str, dt.date]]:
    """[(reporting month YYYY-MM, release date)] from the schedule page.

    Cells are <th>. Years appear only on some cells, so the first year comes from the heading
    'PRESS RELEASE DATES FOR 2026 & 2027'; after that a cell's own year wins, otherwise the year
    carries forward and rolls over when the month goes back past December."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    head = re.search(r"PRESS RELEASE DATES FOR (\d{4})", " ".join(soup.get_text(" ").split()))
    if not head:
        raise SourceError("schedule: 'PRESS RELEASE DATES FOR <year>' heading not found")
    table = None
    for t in soup.find_all("table"):
        first = t.find("tr")
        if first and [" ".join(c.get_text(" ").split()) for c in first.find_all(["th", "td"])] == \
                ["Reporting Period", "Release Date", "Time"]:
            table = t
            break
    if table is None:
        raise SourceError("schedule: table with Reporting Period | Release Date | Time not found")
    names = "|".join(calendar.month_name[1:])
    py = ry = int(head.group(1))
    pm_prev = rm_prev = None
    out = []
    for tr in table.find_all("tr")[1:]:
        cells = [" ".join(c.get_text(" ").split()) for c in tr.find_all(["th", "td"])]
        if not any(cells):
            continue
        p = re.fullmatch(rf"({names})(?: (\d{{4}}))?", cells[0])
        r = re.fullmatch(rf"({names}) (\d{{1,2}})(?:, (\d{{4}}))?", cells[1]) if len(cells) > 1 else None
        if not p or not r:
            raise SourceError(f"schedule: unexpected row {cells}")
        pm, rm = MONTH_NAME[p.group(1)], MONTH_NAME[r.group(1)]
        py = int(p.group(2)) if p.group(2) else py + (1 if pm_prev and pm < pm_prev else 0)
        ry = int(r.group(3)) if r.group(3) else ry + (1 if rm_prev and rm < rm_prev else 0)
        pm_prev, rm_prev = pm, rm
        period, date = f"{py}-{pm:02d}", dt.date(ry, rm, int(r.group(2)))
        lag = (date.year * 12 + date.month) - (py * 12 + pm)
        if not 1 <= lag <= 3:
            raise SourceError(f"schedule: release {date} is {lag} months after {period}")
        out.append((period, date))
    if not out or [p for p, _ in out] != sorted({p for p, _ in out}):
        raise SourceError("schedule: reporting periods not in order")
    return out


def ts_header(h) -> str:
    """'Total\\n_x000D_Construction1' -> 'Total Construction'; drops the footnote digit."""
    s = " ".join(str(h).replace("_x000D_", " ").split())
    return re.sub(r"(?<=[A-Za-z)])\d$", "", s)


TS_COLUMNS = {  # file -> {field: exact header after ts_header}
    "total": {"total": "Total Construction", "x_private": "Total Private Construction",
              "x_res": "Private Residential", "x_nonres": "Private Nonresidential",
              "x_mfg": "Private Manufacturing", "x_pub": "Total Public Construction"},
    "private": {"x_private": "Total Private Construction", "res": "Residential (inc. Improvements)",
                "nonres": "Nonresidential", "mfg": "Manufacturing"},
    "public": {"pub": "Total Public Construction"},
}
TS_TITLES = {"total": "Value of Construction Put in Place - Seasonally Adjusted Annual Rate",
             "private": "Value of Private Construction Put in Place - Seasonally Adjusted Annual Rate",
             "public": "Value of Public Construction Put in Place - Seasonally Adjusted Annual Rate"}


def ts_month(cell) -> str | None:
    """'Aug-26p' -> '2026-08'."""
    m = re.fullmatch(r"([A-Z][a-z]{2})-(\d{2})[pr]?", str(cell).strip()) if isinstance(cell, str) else None
    if not m or m.group(1) not in MONTH_ABBR:
        return None
    yy = int(m.group(2))
    return f"{(1900 if yy >= 50 else 2000) + yy}-{MONTH_ABBR[m.group(1)]:02d}"


def timeseries_from_rows(sheets: dict[str, list[list]]) -> dict[str, dict[str, int]]:
    """{month: {total,res,nonres,pub,mfg}} from the three sheets' rows: title in A1, headers in the
    row whose A cell is 'Date', then one row per month, newest first. totsatime also carries the
    private and public columns; they must equal privsatime and pubsatime."""
    per = {}
    for name, cols in TS_COLUMNS.items():
        rows = sheets[name]
        if not rows or str(rows[0][0]).strip() != TS_TITLES[name]:
            raise SourceError(f"time series {name}: title not found in A1")
        hi = next((i for i, r in enumerate(rows) if r and r[0] == "Date"), None)
        if hi is None:
            raise SourceError(f"time series {name}: 'Date' header row not found")
        heads = [ts_header(h) if h is not None else None for h in rows[hi]]
        idx = {}
        for f, label in cols.items():
            hits = [i for i, h in enumerate(heads) if h == label]
            if len(hits) != 1:
                raise SourceError(f"time series {name}: header {label!r} found {len(hits)} times")
            idx[f] = hits[0]
        got = {}
        for r in rows[hi + 1:]:
            mo = ts_month(r[0]) if r else None
            if mo is None:
                continue
            vals = {f: r[i] if i < len(r) else None for f, i in idx.items()}
            if all(isinstance(v, int) for v in vals.values()):
                if mo in got:
                    raise SourceError(f"time series {name}: {mo} appears twice")
                got[mo] = vals
        if not got:
            raise SourceError(f"time series {name}: no monthly rows")
        per[name] = got
    newest = {n: max(g) for n, g in per.items()}
    if len(set(newest.values())) != 1:
        raise SourceError(f"time series files end in different months: {newest}")
    out = {}
    for mo in sorted(set(per["total"]) & set(per["private"]) & set(per["public"])):
        t, pv, pb = per["total"][mo], per["private"][mo], per["public"][mo]
        if (t["x_private"], t["x_res"], t["x_nonres"], t["x_mfg"], t["x_pub"]) != \
                (pv["x_private"], pv["res"], pv["nonres"], pv["mfg"], pb["pub"]):
            raise SourceError(f"time series: totsatime and privsatime/pubsatime differ for {mo}")
        out[mo] = {"total": t["total"], "res": pv["res"], "nonres": pv["nonres"], "pub": pb["pub"], "mfg": pv["mfg"]}
    return out


def read_timeseries(files: dict[str, bytes]) -> dict[str, dict[str, int]]:
    """{month: {total,res,nonres,pub,mfg}} from totsatime/privsatime/pubsatime."""
    import io
    import openpyxl
    sheets = {}
    for name, content in files.items():
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets[name] = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        except Exception as e:
            raise SourceError(f"time series {name}: {e}")
    return timeseries_from_rows(sheets)


def parse_fred(text: str) -> dict[str, int]:
    """FRED .txt: header lines, then 'DATE VALUE' rows like '2026-08-01 2203129'."""
    out = {}
    for line in text.splitlines():
        m = re.fullmatch(r"\s*(\d{4}-\d{2})-01\s+(-?\d+(?:\.\d+)?)\s*", line)
        if m:
            out[m.group(1)] = round(float(m.group(2)))
    if not out:
        raise SourceError("FRED: no data rows")
    return out


def next_release(schedule: list[tuple[str, dt.date]], data_month: str) -> str:
    want = month_add(data_month, 1)
    for period, date in schedule:
        if period == want:
            return f"{fmt_date(date)} ({calendar.month_abbr[int(want[5:])]} data)"
    return ""


# ---------------------------------------------------------------------------
# Building the update

@dataclass
class Revision:
    month: str
    field: str
    old: int
    new: int

    @property
    def pct(self) -> float:
        return pct(self.new, self.old) if self.old else float("inf")

    def __str__(self):
        return f"{self.month} {self.field}: {self.old:,} → {self.new:,} ({fmt_pct(self.pct)})"


@dataclass
class Result:
    data: dict
    notes: list
    month: str
    revisions: list[Revision]
    added: list[str]
    note: dict | None
    report: list[str] = field(default_factory=list)

    @property
    def big_revisions(self) -> list[Revision]:
        return [r for r in self.revisions if abs(r.pct) > BIG_REVISION_PCT]


def first_max(values: list) -> int:
    """Index of the first maximum (the page uses a strict '>' scan)."""
    p = 0
    for i, v in enumerate(values):
        if v > values[p]:
            p = i
    return p


def current_peak(months: list[list]) -> dict:
    tot = [r[1] for r in months]
    pi = first_max(tot)
    return {"peak_month": months[pi][0], "peak": tot[pi], "k": len(tot) - 1 - pi,
            "index": tot[-1] / tot[pi] * 100}


def then_vs_now(months: list[list], hist: dict) -> dict:
    """The page's 'Then vs. now' reading for total spending (same rules as cmpChart in index.html)."""
    out = current_peak(months)
    k = out["k"]
    h = hist["bust2006"]
    arr = h["total"]
    ids = [month_add(h["start"], i) for i in range(len(arr))]
    p = ids.index(h["peak"]) if h.get("peak") in ids else -1
    if p < 0:
        p = first_max(arr)
    pts = [v / arr[p] * 100 for v in arr[p:]]
    back = next((j for j in range(1, len(pts)) if pts[j] >= 100), -1)
    if back > 0:
        pts = pts[:back + 1]
    out.update(bust_label=h.get("label", "bust2006"), bust_peak_month=month_add(h["start"], p),
               bust_value=pts[k] if k < len(pts) else None, bust_back=back if back > 0 else None)
    return out


def build_note(data: dict, prev_months: dict, rt: ReleaseText, t1: Table1, hist: dict,
               revisions: list[Revision], checks_line: str) -> dict:
    months = data["months"]
    L = dict(zip(FIELDS, months[-1][1:]))
    P = dict(zip(FIELDS, months[-2][1:]))
    m = months[-1][0]
    ymo = month_add(m, -12)
    Y = next((dict(zip(FIELDS, r[1:])) for r in months if r[0] == ymo), None)
    ns = lambda c: "" if c.significant else " (not significant)"
    mom = pct(L["total"], P["total"])
    yoy = pct(L["total"], Y["total"]) if Y else None
    yoy_s = f", {fmt_pct(yoy)} y/y" if yoy is not None else ""
    ytd = data["latest"]["ytd"]
    headline = f"Total {fmt_b(L['total'])}, {fmt_pct(mom)} m/m{ns(rt.total)}{yoy_s}"

    p1 = (f"Total {fmt_b(L['total'])}, {fmt_pct(mom)} m/m{ns(rt.total)}{yoy_s}. "
          f"Year to date ({ytd['label']}): ${ytd['cur']:,.1f}B vs ${ytd['prior']:,.1f}B "
          f"({fmt_pct(pct(ytd['cur'], ytd['prior']))}). Nominal dollars.")
    p2 = (f"Private residential {fmt_b(L['res'])}, {fmt_pct(pct(L['res'], P['res']))} m/m{ns(rt.res)}; "
          f"private nonresidential {fmt_b(L['nonres'])}, {fmt_pct(pct(L['nonres'], P['nonres']))} m/m{ns(rt.nonres)}; "
          f"public {fmt_b(L['pub'])}, {fmt_pct(pct(L['pub'], P['pub']))} m/m{ns(rt.pub)}.")
    sec = {s["name"]: s for s in data["latest"]["sectors"]}
    p3 = "Year over year: " + ", ".join(
        f"{n} {fmt_pct(pct(sec[n]['cur'], sec[n]['yago']))}"
        for n in ("Office (incl. data centers)", "Power", "Manufacturing")) + "."

    parts = []
    for pm in (month_add(m, -1), month_add(m, -2)):
        new = next(r[1] for r in months if r[0] == pm)
        old = prev_months.get(pm, [None, None])[1]
        if old is None or old == new:
            parts.append(f"{month_abbr(pm)} unrevised at {fmt_b(new)}")
        else:
            parts.append(f"{month_abbr(pm)} {fmt_signed_b(new - old)} to {fmt_b(new)}")
    others = {r.month for r in revisions} - {month_add(m, -1), month_add(m, -2)}
    p4 = f"Revisions to total: {'; '.join(parts)}. Other months revised: {len(others)}."

    tv = then_vs_now(months, hist)
    if tv["k"] == 0:
        p5 = f"Then vs. now: total spending is at a new peak ({fmt_b(tv['peak'])})."
    else:
        then = (f"the {tv['bust_label']} was at {tv['bust_value']:.1f} at the same point after its "
                f"{month_abbr(tv['bust_peak_month'])} peak" if tv["bust_value"] is not None
                else f"the {tv['bust_label']} has no data that far out")
        p5 = (f"Then vs. now: {tv['k']} months after the {month_abbr(tv['peak_month'])} peak, total "
              f"spending is at {tv['index']:.1f} (peak = 100); {then}.")
    return {"id": m, "released": data["latest"]["released"], "headline": headline,
            "points": [p1, p2, p3, p4, p5, checks_line], "auto": True}


def check_months(months: list[list]) -> None:
    """Check 7."""
    for r in months:
        if len(r) != 6 or not isinstance(r[0], str) or not all(isinstance(v, int) for v in r[1:]):
            raise CheckFailed(f"check 7: bad row {r[:1]}")
    ids = [r[0] for r in months]
    if ids[0] != "2022-01":
        raise CheckFailed(f"check 7: series starts at {ids[0]}, expected 2022-01")
    for a, b in zip(ids, ids[1:]):
        if b != month_add(a, 1):
            raise CheckFailed(f"check 7: {a} is followed by {b} (gap, duplicate or out of order)")


def check_ties(months: list[list], which: set[str]) -> None:
    """Check 2 for the DATA rows."""
    for r in months:
        if r[0] in which:
            _, total, res, nonres, pub, _ = r
            if abs(res + nonres + pub - total) > TIE_TOLERANCE:
                raise CheckFailed(f"check 2: {r[0]} res+nonres+pub = {res + nonres + pub:,} vs total {total:,}")


def check_table1(t1: Table1) -> None:
    """Check 2 for Table 1: private residential + nonresidential = total private."""
    for i, mo in enumerate(t1.cols):
        res, nonres, tp = (t1.values[k][i] for k in ("res", "nonres", "private_total"))
        if None in (res, nonres, tp):
            continue
        if abs(res + nonres - tp) > TIE_TOLERANCE:
            raise CheckFailed(f"check 2: {mo} private res+nonres = {res + nonres:,} vs total private {tp:,}")


def check_against_text(t1: Table1, rt: ReleaseText) -> None:
    """Checks 1 and 3: Table 1 agrees with the release text."""
    if t1.month != rt.month:
        raise CheckFailed(f"check 1: Table 1 month {t1.month} differs from the release page month {rt.month}")
    m, pm = t1.cols[0], t1.cols[1]
    for key, c in (("total", rt.total), ("res", rt.res), ("nonres", rt.nonres), ("pub", rt.pub)):
        cur, prev = t1.get(key, m), t1.get(key, pm)
        if abs(cur / 1000 - c.level_b) > 0.05 + 1e-9:
            raise CheckFailed(f"check 3: {key} Table 1 {cur / 1000:.1f}B vs text {c.level_b}B")
        if abs(pct(cur, prev) - c.pct) > 0.1 + 1e-9:
            raise CheckFailed(f"check 3: {key} m/m Table 1 {pct(cur, prev):.2f}% vs text {c.pct}%")
    yoy = pct(t1.get("total", m), t1.get("total", t1.yago))
    if abs(yoy - rt.yoy.pct) > 0.1 + 1e-9:
        raise CheckFailed(f"check 3: total y/y Table 1 {yoy:.2f}% vs text {rt.yoy.pct}%")


def build_update(data: dict, notes: list, hist: dict, t1: Table1, rt: ReleaseText | None,
                 ts: dict | None, schedule: list | None, fred: dict | None) -> Result | None:
    """Return the new DATA and NOTES, or None when Table 1 has nothing newer than the page."""
    latest = data["latest"]
    if t1.month <= latest["month"]:
        return None
    report = []
    check_table1(t1)
    if rt is not None:
        check_against_text(t1, rt)

    old = {r[0]: r for r in data["months"]}
    rows = {k: list(v) for k, v in old.items()}
    agreed_ts = []
    if ts is not None:
        for mo in old:
            if mo in ts:
                rows[mo] = [mo] + [ts[mo][f] for f in FIELDS]
    for i, mo in enumerate(t1.cols):
        if i == len(t1.cols) - 1 and mo not in rows:
            continue  # year-ago month before 2022-01
        new = [mo] + [t1.values[f][i] for f in FIELDS]
        if ts is not None and mo in ts:
            if [ts[mo][f] for f in FIELDS] == new[1:]:
                agreed_ts.append(mo)
            else:
                report.append(f"Time series and Table 1 differ for {mo}; Table 1 used.")
        rows[mo] = new
    months = [rows[k] for k in sorted(rows)]
    check_months(months)

    revisions = [Revision(mo, f, old[mo][j + 1], rows[mo][j + 1])
                 for mo in sorted(old) for j, f in enumerate(FIELDS) if old[mo][j + 1] != rows[mo][j + 1]]
    added = [r[0] for r in months if r[0] not in old]
    check_ties(months, {r.month for r in revisions} | set(added))

    m = t1.month
    def sector(name, key, group=False):
        s = {"name": name}
        if group:
            s["group"] = True
        s.update(cur=t1.get(key, m), prev=t1.get(key, t1.cols[1]), yago=t1.get(key, t1.yago))
        if any(not isinstance(s[k], int) for k in ("cur", "prev", "yago")):
            raise CheckFailed(f"Table 1: sector {name} is missing a value")
        return s
    sectors = [sector("Private nonresidential", "nonres", True), sector("Manufacturing", "mfg"),
               sector("Power", "power"), sector("Office (incl. data centers)", "office"),
               sector("Commercial", "commercial"), sector("Public", "pub", True),
               sector("Highway and street", "highway"), sector("Educational", "educational")]

    sched_entry = dict(schedule or [])
    if rt is not None:
        released = fmt_date(rt.released)
        moe = {"total": rt.total.significant, "res": rt.res.significant,
               "nonres": rt.nonres.significant, "pub": rt.pub.significant}
        if rt.ytd_months != int(m[5:]):
            raise CheckFailed(f"YTD covers {rt.ytd_months} months but the data month is {m}")
        mabbr = calendar.month_abbr[int(m[5:])]
        ytd = {"cur": round(rt.ytd_cur_b, 1), "prior": round(rt.ytd_prior_b, 1),
               "label": f"Jan{EN_DASH}{mabbr} {m[:4]} vs Jan{EN_DASH}{mabbr} {int(m[:4]) - 1}"}
    else:
        released = fmt_date(sched_entry[m]) if m in sched_entry else ""
        moe, ytd = latest.get("moe"), latest.get("ytd")
        report.append("Release text not read: significance flags and year to date carried over "
                      f"from {month_abbr(latest['month'])}; no automatic note written.")
    nxt = next_release(schedule, m) if schedule else ""
    if not nxt:
        report.append("Next release date not found in the schedule; left blank.")

    new_latest = {"month": m, "status": "preliminary", "released": released, "next": nxt,
                  "moe": moe, "ytd": ytd,
                  "sectorNote": f"Selected categories from the {month_full(m)} release. "
                                "Census counts data centers under Office.",
                  "sectors": sectors}
    new_data = {**data, "months": months, "latest": {**latest, **new_latest}}

    # Checks line for the note: which sources agreed, what was single source.
    agreed = ["Table 1 and the release text agreed on total, residential, nonresidential and public "
              "levels and changes"]
    if ts is not None:
        agreed.append(f"the monthly time series matched Table 1 for {len(agreed_ts)} of "
                      f"{sum(1 for c in t1.cols if c in ts)} months")
    else:
        agreed.append("monthly time series not read")
        report.append("Time series files not read: only the Table 1 months were updated.")
    fred_note = "FRED not reached"
    if fred:
        fm = {f: fred[f].get(m) for f in FIELDS if f in fred}
        if all(v is None for v in fm.values()):
            fred_note = "FRED not yet updated"
        else:
            bad = [f for f, v in fm.items() if v is not None and v != t1.get(f, m)]
            fred_note = "FRED agreed" if not bad else f"FRED differed on {', '.join(bad)}"
    report.append(f"FRED: {fred_note}.")
    checks_line = (f"Checks: {'; '.join(agreed)}; {fred_note}. "
                   "Single source: sector rows (Table 1 only).")

    new_notes = [n for n in notes if n.get("id") != m]
    note = None
    if rt is not None:
        note = build_note(new_data, old, rt, t1, hist, revisions, checks_line)
        new_notes.append(note)
    return Result(new_data, new_notes, m, revisions, added, note, report)


def build_summary(data: dict) -> dict:
    months = data["months"]
    L, P = months[-1], months[-2]
    Y = next((r for r in months if r[0] == month_add(L[0], -12)), None)
    lt = data["latest"]
    tv = current_peak(months)
    return {"month": lt["month"], "status": lt["status"], "released": lt["released"], "next": lt["next"],
            "total": L[1], "mom": round(pct(L[1], P[1]), 1),
            "yoy": round(pct(L[1], Y[1]), 1) if Y else None,
            "ytd_cur": lt["ytd"]["cur"], "ytd_prior": lt["ytd"]["prior"],
            "res": L[2], "nonres": L[3], "pub": L[4],
            "peak_month": tv["peak_month"], "pct_below_peak": round(100 - tv["index"], 1)}


def commit_message(res: Result) -> str:
    lines = [f"Data: {month_abbr(res.month)} construction spending (preliminary)", ""]
    if res.added:
        lines.append("Added months: " + ", ".join(res.added))
    if res.revisions:
        lines.append(f"Revised values ({len(res.revisions)}):")
        lines += [f"- {r}" for r in res.revisions]
    else:
        lines.append("No revisions to stored months.")
    if res.big_revisions:
        lines += ["", f"Moves over {BIG_REVISION_PCT:.0f}% (please review):"]
        lines += [f"- {r}" for r in res.big_revisions]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Running

class Report:
    def __init__(self):
        self.lines: list[str] = []

    def __call__(self, line: str = ""):
        self.lines.append(line)
        print(line)

    def write(self):
        path = os.environ.get("GITHUB_STEP_SUMMARY")
        if path:
            with open(path, "a", encoding="utf-8") as f:
                f.write("\n".join(self.lines) + "\n")


def set_output(name: str, value: str) -> None:
    path = os.environ.get("GITHUB_OUTPUT")
    if path:
        with open(path, "a", encoding="utf-8") as f:
            f.write(f"{name}={value}\n")


def get_table1(say) -> Table1:
    for url, reader in ((URL_TABLE1_XLSX, read_table1_xlsx), (URL_TABLE1_PDF, read_table1_pdf)):
        try:
            t1 = reader(fetch(url))
            t1.source = url.rsplit("/", 1)[-1]
            say(f"- Table 1 read from {t1.source}: newest month {t1.month}.")
            return t1
        except SourceError as e:
            say(f"- Table 1 from {url.rsplit('/', 1)[-1]} failed: {e}")
    raise CheckFailed("Table 1 could not be read from release.xlsx or release.pdf")


def optional(say, what, fn):
    try:
        return fn()
    except SourceError as e:
        say(f"- {what} not read: {e}")
        return None


def run(page: Path, summary_path: Path, dry_run: bool, commit_msg: Path | None) -> int:
    say = Report()
    say(f"## Construction spending update{' (dry run)' if dry_run else ''}")
    set_output("changed", "false")
    set_output("big_revisions", "false")
    try:
        html = page.read_text(encoding="utf-8")
        data, notes, hist_line = load_page(html)
        hist = parse_block(hist_line)
        say(f"- Page has data through {data['latest']['month']}.")
        t1 = get_table1(say)
        if t1.month <= data["latest"]["month"]:
            say(f"- Nothing new: Table 1 newest month is {t1.month}.")
            return 0
        rt = optional(say, "Release text", lambda: parse_release_text(fetch(URL_RELEASE_TEXT)))
        ts = optional(say, "Time series",
                      lambda: read_timeseries({k: fetch(u) for k, u in URL_TIMESERIES.items()}))
        schedule = optional(say, "Release schedule", lambda: parse_schedule(fetch(URL_SCHEDULE)))
        fred = {}
        for f, sid in FRED_SERIES.items():
            try:
                fred[f] = parse_fred(fetch_once(URL_FRED.format(sid)).decode("utf-8", "replace"))
            except SourceError:
                fred = None
                break

        res = build_update(data, notes, hist, t1, rt, ts, schedule, fred)
        new_html = render_page(html, res.data, res.notes)
        verify_page(html, new_html)
        if load_page(new_html)[2] != hist_line:
            raise CheckFailed("check 6: HIST changed")
    except CheckFailed as e:
        say(f"\n**Check failed, nothing written:** {e}")
        say.write()
        return 1

    say("\nAll checks passed.\n")
    for line in res.report:
        say(f"- {line}")
    say(f"\n### Changes for {month_abbr(res.month)}")
    if res.added:
        say("- Added months: " + ", ".join(res.added))
    for r in res.revisions:
        say(f"- Revised {r}")
    if not res.revisions:
        say("- No revisions to stored months.")
    for k in ("released", "next", "moe", "ytd", "sectorNote", "sectors"):
        o, n = data["latest"].get(k), res.data["latest"].get(k)
        if o != n:
            say(f"- latest.{k}: `{dump(o)}` → `{dump(n)}`")
    if res.note:
        say("\n### Automatic note\n")
        say(f"**{res.note['headline']}**\n")
        for p in res.note["points"]:
            say(f"- {p}")
    if res.big_revisions:
        say(f"\n**{len(res.big_revisions)} value(s) moved more than {BIG_REVISION_PCT:.0f}%.**")

    if not dry_run:
        page.write_text(new_html, encoding="utf-8")
        summary_path.parent.mkdir(parents=True, exist_ok=True)
        summary_path.write_text(json.dumps(build_summary(res.data), ensure_ascii=False, indent=2) + "\n",
                                encoding="utf-8")
        if commit_msg:
            commit_msg.write_text(commit_message(res), encoding="utf-8")
        set_output("changed", "true")
        set_output("big_revisions", "true" if res.big_revisions else "false")
    else:
        say("\nDry run: nothing written.")
    say.write()
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--dry-run", action="store_true", help="fetch, parse and check; write nothing")
    ap.add_argument("--page", type=Path, default=ROOT / "index.html")
    ap.add_argument("--summary", type=Path, default=ROOT / "data" / "summary.json")
    ap.add_argument("--commit-msg", type=Path, help="where to write the commit message")
    ap.add_argument("--summary-only", action="store_true",
                    help="rewrite data/summary.json from the page and exit (no fetching)")
    a = ap.parse_args(argv)
    if a.summary_only:
        data = load_page(a.page.read_text(encoding="utf-8"))[0]
        a.summary.parent.mkdir(parents=True, exist_ok=True)
        a.summary.write_text(json.dumps(build_summary(data), ensure_ascii=False, indent=2) + "\n",
                             encoding="utf-8")
        return 0
    return run(a.page, a.summary, a.dry_run, a.commit_msg)


if __name__ == "__main__":
    sys.exit(main())
